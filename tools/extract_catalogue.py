"""Extract a private DOCX catalogue; never assume a heading is executable regex.

Usage: python tools/extract_catalogue.py input.docx --overrides overrides.json
       --output private-catalogue/catalogue.json
Review overrides are keyed by zero-based Word paragraph index. The generated
catalogue contains full technical source text and MUST retain source access limits.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path
from docx import Document


def clean(text):
    return re.sub(r"[\t \xa0]+", " ", text.replace("\u200b", "").replace("\ufeff", "")).strip()


def heading_pattern(title):
    title = re.sub(r"^.*?Syslog message\s*[:\-]\s*", "", title, flags=re.I)
    # Only the explicitly used .* convention is executable; punctuation is literal.
    return ".*".join(re.escape(clean(x)).replace(r"\ ", r"\s+") for x in title.split(".*"))


def extract(path, overrides):
    checksum = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    if overrides.get("source_sha256") and overrides["source_sha256"] != checksum:
        raise ValueError("Source changed: review paragraph overrides before updating source_sha256")
    doc = Document(path)
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        # Include hyperlink text even with older python-docx versions.
        text = "".join(p._p.xpath(".//w:t/text()"))
        # p.text retains explicit line breaks; use it when no hyperlink text was lost.
        if len(p.text.replace("\n", "")) >= len(text):
            text = p.text
        paragraphs.append({"index": i, "style": p.style.name, "text": clean(text)})
    starts = [p["index"] for p in paragraphs if (
        p["style"] == "Heading 2" and p["text"] and not p["text"].startswith('The "')
    ) or str(p["index"]) in overrides.get("extra_sections", [])]
    boundaries = sorted(set(starts + [p["index"] for p in paragraphs if p["style"] == "Heading 1"] + [len(paragraphs)]))
    rules = {}
    audit = []
    for start in starts:
        end = next(i for i in boundaries if i > start)
        title = paragraphs[start]["text"]
        override = overrides.get("sections", {}).get(str(start), {})
        end = override.get("paragraph_end_exclusive", end)
        source = paragraphs[start:end]
        pattern = override.get("pattern", heading_pattern(title))
        re.compile(pattern)
        family = next((p["text"] for p in reversed(paragraphs[:start]) if p["style"] == "Heading 1"), "Unknown")
        body = "\n\n".join(p["text"] for p in source[1:] if p["text"])
        marker = next((j for j, p in enumerate(source) if re.match(r"^(Actions?|Recommended Action|Solution)\s*:?$", p["text"], re.I)), None)
        actions = "\n\n".join(p["text"] for p in source[marker + 1:] if p["text"]) if marker is not None else body
        refs = []
        for p in doc.paragraphs[start:end]:
            for link in p._p.xpath(".//w:hyperlink"):
                rid = link.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                if rid and rid in doc.part.rels:
                    refs.append(doc.part.rels[rid].target_ref)
        samples = [clean(line) for p in source[1:] for line in p["text"].splitlines()
                   if re.match(r"^(?:\[)?(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+\d|^<host>|^fpc\d|^\d{4}-\d\d-\d\d|^(?:alarmd|craftd)\[", clean(line), re.I)]
        matched = [x for x in samples if re.search(pattern, x, re.I)]
        enabled = override.get("enabled", bool(matched))
        if not matched and enabled and not override.get("positive_examples"):
            raise ValueError(f"Section {start} enabled without positive evidence")
        pattern_id = "jls-" + hashlib.sha256(pattern.casefold().encode()).hexdigest()[:16]
        item = rules.setdefault(pattern_id, {
            "id": pattern_id, "title": title, "pattern": pattern, "enabled": enabled,
            "families": [], "kb_references": [], "pr_references": [], "links": [],
            "sources": [], "positive_examples": [], "negative_examples": [],
            "review_notes": [], "version_notes": [], "platform_notes": [],
        })
        item["sources"].append({"paragraph_start": start, "paragraph_end": end - 1,
                                "title": title, "technical_text": body, "actions": actions})
        item["families"].append(family)
        item["kb_references"] += re.findall(r"\bKB\d+", body, re.I)
        item["pr_references"] += ["PR" + x for x in re.findall(r"\bPR\s*/?\s*(\d+)", body, re.I)]
        item["links"] += refs
        item["positive_examples"] += override.get("positive_examples", matched)
        item["negative_examples"] += override.get("negative_examples", [])
        item["review_notes"] += override.get("notes", [])
        item["version_notes"] += [p["text"] for p in source if re.search(r"\b\d{2}\.\d[RF]\d|\bJunos\s+(?:OS\s+)?\d{2}\.\d", p["text"], re.I)]
        item["platform_notes"] += [p["text"] for p in source if re.search(r"\b(?:MPC\d|MX\d|T4000|routing.engine)", p["text"], re.I)]
        audit.append({"paragraph": start, "id": pattern_id, "title": title, "sample_count": len(samples),
                      "matched_source_samples": len(matched), "enabled": enabled, "override": bool(override)})
    for item in rules.values():
        for key in ("families", "kb_references", "pr_references", "links", "positive_examples", "negative_examples", "review_notes", "version_notes", "platform_notes"):
            item[key] = list(dict.fromkeys(item[key]))
    return {
        "schema_version": 1, "catalogue_id": overrides["catalogue_id"], "version": overrides["version"],
        "source": {"filename": Path(path).name, "sha256": checksum,
                   "revision_date": overrides.get("revision_date"), "classification": "confidential",
                   "notice": next((p["text"] for p in paragraphs if "trade secrets" in p["text"]), ""),
                   "general_guidance": [p for p in paragraphs if 24 <= p["index"] < min(starts) and p["text"]],
                   "references": [p["text"] for p in paragraphs if p["index"] > max(starts) and p["text"].startswith("http")]},
        "rules": list(rules.values()), "extraction_audit": audit,
        "supplementary_rules": overrides.get("supplementary_rules", []),
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("document")
    parser.add_argument("--overrides", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    catalogue = extract(args.document, json.loads(Path(args.overrides).read_text()))
    catalogue["rules"].extend(catalogue.pop("supplementary_rules"))
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(catalogue, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"sections": len(catalogue["extraction_audit"]), "rules": len(catalogue["rules"]),
                      "enabled": sum(r["enabled"] for r in catalogue["rules"])}))
